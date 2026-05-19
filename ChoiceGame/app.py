from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path
from typing import Literal
from datetime import datetime

try:
    import openpyxl
except ImportError:
    openpyxl = None

from flask import (
    Flask,
    Response,
    abort,
    redirect,
    render_template,
    request,
    session,
    url_for,
)


BASE_DIR = Path(__file__).resolve().parent


Condition = Literal["Guilty", "Innocent"]
Case = Literal["Arson", "Theft"]
Choice = Literal["A", "B"]

def log_to_excel(pid: str, case: str, condition: str, step_index: int, video: str, choice: str):
    if openpyxl is None:
        print("Warning: openpyxl is not installed. Skip logging to Excel.")
        return
        
    file_path = BASE_DIR / "results.xlsx"
    if not file_path.exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Timestamp", "ParticipantID", "Case", "Condition", "StepIndex", "Video", "Choice"])
    else:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
    
    ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), pid, case, condition, step_index, video, choice])
    wb.save(file_path)

# YouTube mapping placeholders
# - Fill each video's YouTube videoId (not the full URL)
# - Example: https://www.youtube.com/watch?v=dQw4w9WgXcQ  ->  "dQw4w9WgXcQ"
YOUTUBE_VIDEO_IDS: dict[str, str] = {
    "Guilty1.mp4": "bvNTJILT0fE",
    "Guilty2-1.mp4": "3nravlMn0wo",
    "Guilty2-2.mp4": "NrB6rORbqH0",
    "Guilty3.mp4": "nJ5j5tju3Zg",
    "Guilty4-1.mp4": "9qccHF6P_eI",
    "Guilty4-2.mp4": "j2x_5JCYy5I",
    "Guilty5.mp4": "ZXOLG9Bbx6s",
    "Guilty6-1.mp4": "xcDC6VXubkg",
    "Guilty6-2.mp4": "-4ggtpRgecw",
    "Guilty7.mp4": "Fl8nqdlsgwI",
    "Innocent1.mp4": "Fb4-YamV45E",
    "Innocent2-1.mp4": "7vD-KzJ9UNM",
    "Innocent2-2.mp4": "UQp13fRa0ig",
    "Innocent3.mp4": "2_aAHkoipF0",
    "Innocent4-1.mp4": "ny9820iXPEw",
    "Innocent4-2.mp4": "gr5NRIy73Wg",
    "Innocent5.mp4": "rEMvZCmGK00",
    "Innocent6-1.mp4": "OIgEbhY0UYo",
    "Innocent6-2.mp4": "2g_nL2xDMt0",
    "Innocent7.mp4": "n4UH9zKVPpg",
    "Theft_Guilty1.mp4": "U-4P5l70AIA",
    "Theft_Guilty2-1.mp4": "FRSSkfDA60Q",
    "Theft_Guilty2-2.mp4": "ppST6j4VvcA",
    "Theft_Guilty3.mp4": "jk4qOaqd7cA",
    "Theft_Guilty4-1.mp4": "iQpfr4Uz2fs",
    "Theft_Guilty4-2.mp4": "GxAJxGwAgKQ",
    "Theft_Guilty5.mp4": "l9WQ2drRW6c",
    "Theft_Guilty6-1.mp4": "TPr7a9K1uqI",
    "Theft_Guilty6-2.mp4": "Lh5BiQLFFIg",
    "Theft_Guilty7.mp4": "WfUq0jigYSc",
    "Theft_Innocent1.mp4": "M5rQ4QDneCE",
    "Theft_Innocent2-1.mp4": "PFiV5SPVRSM",
    "Theft_Innocent2-2.mp4": "XwUnuOFTcHA",
    "Theft_Innocent3.mp4": "ciSlIwU3_kM",
    "Theft_Innocent4-1.mp4": "FYLant-Ai9A",
    "Theft_Innocent4-2.mp4": "P9w5PmF6RjI",
    "Theft_Innocent5.mp4": "LcpU9L4ki9U",
    "Theft_Innocent6-1.mp4": "Ucolst6-M3I",
    "Theft_Innocent6-2.mp4": "Uk5Vh5DFHhc",
    "Theft_Innocent7.mp4": "AMU235ytfZk",
}


@dataclass(frozen=True)
class Step:
    video: str
    question: str | None = None
    a_label: str | None = None
    b_label: str | None = None
    next_default: int | None = None
    next_if_a: int | None = None
    next_if_b: int | None = None

    @property
    def has_choice(self) -> bool:
        return self.question is not None


def build_timeline(case: Case, condition: Condition) -> list[Step]:
    if case == "Theft":
        if condition == "Guilty":
            return [
                Step(
                    video="Theft_Guilty1.mp4",
                    question="Please choose:",
                    a_label="A) Buy a latte",
                    b_label="B) Buy milk",
                    next_if_a=1,  # Theft_Guilty2-1
                    next_if_b=2,  # Theft_Guilty2-2
                ),
                Step(video="Theft_Guilty2-1.mp4", next_default=3),
                Step(video="Theft_Guilty2-2.mp4", next_default=3),
                Step(
                    video="Theft_Guilty3.mp4",
                    question="Please choose:",
                    a_label="A) Walk quickly across the square",
                    b_label="B) Walk across the square at a normal pace",
                    next_if_a=4,  # Theft_Guilty4-1
                    next_if_b=5,  # Theft_Guilty4-2
                ),
                Step(video="Theft_Guilty4-1.mp4", next_default=6),
                Step(video="Theft_Guilty4-2.mp4", next_default=6),
                Step(
                    video="Theft_Guilty5.mp4",
                    question="Please choose:",
                    a_label="A) Put it in the backpack",
                    b_label="B) Put it in the pocket",
                    next_if_a=7,  # Theft_Guilty6-1
                    next_if_b=8,  # Theft_Guilty6-2
                ),
                Step(video="Theft_Guilty6-1.mp4", next_default=9),
                Step(video="Theft_Guilty6-2.mp4", next_default=9),
                Step(video="Theft_Guilty7.mp4", next_default=10),
            ]

        # Theft Innocent
        return [
            Step(
                video="Theft_Innocent1.mp4",
                question="Please choose:",
                a_label="A) Buy a latte",
                b_label="B) Buy milk",
                next_if_a=1,  # Theft_Innocent2-1
                next_if_b=2,  # Theft_Innocent2-2
            ),
            Step(video="Theft_Innocent2-1.mp4", next_default=3),
            Step(video="Theft_Innocent2-2.mp4", next_default=3),
            Step(
                video="Theft_Innocent3.mp4",
                question="Please choose:",
                a_label="A) Take a picture of the sea",
                b_label="B) Take a picture of the square",
                next_if_a=4,  # Theft_Innocent4-1
                next_if_b=5,  # Theft_Innocent4-2
            ),
            Step(video="Theft_Innocent4-1.mp4", next_default=6),
            Step(video="Theft_Innocent4-2.mp4", next_default=6),
            Step(
                video="Theft_Innocent5.mp4",
                question="Please choose:",
                a_label="A) Look at the cruise ship on the left",
                b_label="B) Look at the cruise ship on the right",
                next_if_a=7,  # Theft_Innocent6-1
                next_if_b=8,  # Theft_Innocent6-2
            ),
            Step(video="Theft_Innocent6-1.mp4", next_default=9),
            Step(video="Theft_Innocent6-2.mp4", next_default=9),
            Step(video="Theft_Innocent7.mp4", next_default=10),
        ]

    # Arson
    if condition == "Guilty":
        return [
            Step(
                video="Guilty1.mp4",
                question="Please choose:",
                a_label="A) Park on a small road 400 meters away, then walk over",
                b_label="B) Park in a public parking lot",
                next_if_a=1,  # Guilty2-1
                next_if_b=2,  # Guilty2-2
            ),
            Step(video="Guilty2-1.mp4", next_default=3),
            Step(video="Guilty2-2.mp4", next_default=3),
            Step(
                video="Guilty3.mp4",
                question="Please choose:",
                a_label="A) Carefully pour gasoline on the load-bearing pillar",
                b_label="B) Quickly pour gasoline onto the ground",
                next_if_a=4,  # Guilty4-1
                next_if_b=5,  # Guilty4-2
            ),
            Step(video="Guilty4-1.mp4", next_default=6),
            Step(video="Guilty4-2.mp4", next_default=6),
            Step(
                video="Guilty5.mp4",
                question="Please choose:",
                a_label="A) Drive home via the main road",
                b_label="B) Drive home via the side road",
                next_if_a=7,  # Guilty6-1
                next_if_b=8,  # Guilty6-2
            ),
            Step(video="Guilty6-1.mp4", next_default=9),
            Step(video="Guilty6-2.mp4", next_default=9),
            Step(video="Guilty7.mp4", next_default=10),
        ]

    # Innocent
    return [
        Step(
            video="Innocent1.mp4",
            question="Please choose:",
            a_label="A) Watch a cartoon movie",
            b_label="B) Watch an action movie",
            next_if_a=1,  # Innocent2-1
            next_if_b=2,  # Innocent2-2
        ),
        Step(video="Innocent2-1.mp4", next_default=3),
        Step(video="Innocent2-2.mp4", next_default=3),
        Step(
            video="Innocent3.mp4",
            question="Please choose:",
            a_label="A) Listen to soft music",
            b_label="B) Listen to upbeat music",
            next_if_a=4,  # Innocent4-1
            next_if_b=5,  # Innocent4-2
        ),
        Step(video="Innocent4-1.mp4", next_default=6),
        Step(video="Innocent4-2.mp4", next_default=6),
        Step(
            video="Innocent5.mp4",
            question="Please choose:",
            a_label="A) Drive home via the main road",
            b_label="B) Drive home via the side road",
            next_if_a=7,  # Innocent6-1
            next_if_b=8,  # Innocent6-2
        ),
        Step(video="Innocent6-1.mp4", next_default=9),
        Step(video="Innocent6-2.mp4", next_default=9),
        Step(video="Innocent7.mp4", next_default=10),
    ]


def parse_participant_id(raw: str) -> tuple[Case, Condition]:
    parts = raw.split("-")
    if len(parts) != 2:
        raise ValueError("Invalid format")
    case_part, cond_part = parts
    
    if case_part == "1":
        case = "Arson"
    elif case_part == "2":
        case = "Theft"
    else:
        raise ValueError("Unknown case part (must be 1 or 2)")
        
    cond_part = cond_part.strip().lower()
    if cond_part in ["guilty", "1"]:
        condition = "Guilty"
    elif cond_part in ["innocent", "2"]:
        condition = "Innocent"
    elif cond_part.isdigit():
        condition = "Guilty" if int(cond_part) % 2 == 1 else "Innocent"
    else:
        raise ValueError("Unknown condition part")
        
    return case, condition


def _ensure_session_initialized() -> None:
    if "participant_id" not in session or "condition" not in session or "timeline" not in session:
        abort(400, "Session not initialized. Go back to home and start again.")


def youtube_id_for(video_name: str) -> str | None:
    vid = (YOUTUBE_VIDEO_IDS.get(video_name) or "").strip()
    return vid or None

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "dev-secret-key-change-me")
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_HTTPONLY"] = True


@app.get("/")
def index() -> str:
    return render_template("index.html")


@app.post("/start")
def start() -> Response:
    raw = (request.form.get("participant_id") or "").strip()
    try:
        case, condition = parse_participant_id(raw)
    except ValueError:
        return render_template("index.html", error="ID format must be '1-1' or '2-Guilty' (1=Arson, 2=Theft)."), 400

    pid = raw
    timeline = build_timeline(case, condition)

    # Validate YouTube mapping configured
    missing_ids = [s.video for s in timeline if youtube_id_for(s.video) is None]
    if missing_ids:
        return (
            render_template(
                "index.html",
                error="These videos do not have a configured YouTube videoId in app.py (YOUTUBE_VIDEO_IDS): "
                + ", ".join(missing_ids),
            ),
            500,
        )

    session.clear()
    session["participant_id"] = pid
    session["case"] = case
    session["condition"] = condition
    session["timeline"] = [s.__dict__ for s in timeline]
    session["idx"] = 0

    return redirect(url_for("play"))


@app.get("/play")
def play() -> str:
    _ensure_session_initialized()

    idx = int(session["idx"])
    timeline = session["timeline"]
    if idx < 0 or idx >= len(timeline):
        return render_template("done.html")

    step = timeline[idx]
    video = step["video"]
    youtube_id = youtube_id_for(video)
    if youtube_id is None:
        abort(500, f"Missing YouTube mapping for {video}")

    return render_template(
        "play.html",
        participant_id=session["participant_id"],
        youtube_id=youtube_id,
        has_choice=bool(step.get("question")),
        question=step.get("question"),
        a_label=step.get("a_label"),
        b_label=step.get("b_label"),
    )

@app.post("/next")
def next_step() -> Response:
    _ensure_session_initialized()
    idx = int(session["idx"])
    timeline = session["timeline"]
    if idx < 0 or idx >= len(timeline):
        return redirect(url_for("play"))

    step = timeline[idx]
    if step.get("question"):
        abort(400, "This step requires a choice.")

    nxt = step.get("next_default")
    session["idx"] = int(nxt) if nxt is not None else (idx + 1)
    return redirect(url_for("play"))


@app.post("/choice")
def choice() -> Response:
    _ensure_session_initialized()
    idx = int(session["idx"])
    timeline = session["timeline"]
    if idx < 0 or idx >= len(timeline):
        return redirect(url_for("play"))

    step = timeline[idx]
    if not step.get("question"):
        abort(400, "This step does not accept a choice.")

    c = (request.form.get("choice") or "").strip().upper()
    if c not in ("A", "B"):
        abort(400, "Invalid choice.")

    next_idx = step["next_if_a"] if c == "A" else step["next_if_b"]
    if next_idx is None:
        abort(400, "Choice routing not configured.")

    # Log to excel
    try:
        log_to_excel(
            pid=session["participant_id"],
            case=session.get("case", "Unknown"),
            condition=session["condition"],
            step_index=idx,
            video=step["video"],
            choice=c
        )
    except Exception as e:
        print(f"Failed to log to Excel: {e}")

    session["idx"] = int(next_idx)
    return redirect(url_for("play"))


if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000)
